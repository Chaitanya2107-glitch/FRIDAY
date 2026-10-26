"""
services/files/extract.py

Reads the text inside a file, for the file index. Plain text and code are read
directly; PDF, Word, PowerPoint and Excel through their libraries. Everything
else (images, audio, video, programs, game data, logs) isn't read — those files
are still found by name.

Secrets are never read: key files, password databases and .env files are
skipped entirely (see is_private), so they never reach the index or the AI.
"""

import logging
import re
from pathlib import Path

from core.logger import logger

logging.getLogger("pypdf").setLevel(logging.ERROR)  # it warns about every oddly-built PDF

TEXT_EXTENSIONS = {
    # notes and data
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".xml", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".tex", ".rtf", ".html", ".htm", ".srt",
    # code
    ".py", ".ipynb", ".js", ".jsx", ".ts", ".tsx", ".mjs", ".cjs", ".vue", ".svelte", ".css", ".scss",
    ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".go", ".rs", ".rb", ".php", ".swift",
    ".dart", ".lua", ".r", ".m", ".sql", ".sh", ".bat", ".cmd", ".ps1", ".ino", ".asm", ".v", ".vhd",
    ".gradle", ".dockerfile",
}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx"}
READABLE_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS

MAX_TEXT_BYTES = 1_000_000        # of a text/code file; the rest is ignored
MAX_DOCUMENT_BYTES = 30_000_000   # bigger PDFs/Office files aren't opened
MAX_CHARS = 400_000               # of text kept per file
MAX_PDF_PAGES = 300
MAX_SHEET_ROWS = 2000

# Also Office's "~$name.docx" lock files, which are only open-document markers
PRIVATE_NAME = re.compile(
    r"^~\$|^\.env(\..*)?$|^id_(rsa|dsa|ecdsa|ed25519)|\.(pem|key|pfx|p12|kdbx|keystore|jks|ovpn)$"
    r"|password|passwd|secret|credential",
    re.IGNORECASE,
)


def is_private(path: Path) -> bool:
    return bool(PRIVATE_NAME.search(path.name))


def is_readable(path: Path) -> bool:
    return path.suffix.lower() in READABLE_EXTENSIONS and not is_private(path)


def extract_text(path: Path, size: int) -> str | None:
    """The file's text, or None if it has none we can read."""
    ext = path.suffix.lower()
    if not is_readable(path):
        return None
    try:
        if ext in TEXT_EXTENSIONS:
            text = _read_text(path)
        elif size > MAX_DOCUMENT_BYTES:
            return None
        elif ext == ".pdf":
            text = _read_pdf(path)
        elif ext == ".docx":
            text = _read_docx(path)
        elif ext == ".pptx":
            text = _read_pptx(path)
        else:
            text = _read_xlsx(path)
    except Exception as e:  # damaged, encrypted or locked files are common — skip them
        logger.info(f"Couldn't read {path}: {e}")
        return None
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text or "")  # stray control bytes from PDFs
    text = re.sub(r"[ \t]+", " ", text).strip()
    return text[:MAX_CHARS] or None


def _read_text(path: Path) -> str | None:
    with open(path, "rb") as f:
        data = f.read(MAX_TEXT_BYTES)
    if b"\x00" in data[:4096]:
        return None  # binary despite its extension
    return data.decode("utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES])


def _read_docx(path: Path) -> str:
    import docx
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return "\n".join(parts)


def _read_pptx(path: Path) -> str:
    from pptx import Presentation
    slides = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        slides.append(f"Slide {number}: " + "\n".join(texts))
    return "\n\n".join(slides)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True):
                cells = [str(value) for value in row if value is not None]
                if cells:
                    rows.append(" | ".join(cells))
            sheets.append(f"Sheet {sheet.title}:\n" + "\n".join(rows))
        return "\n\n".join(sheets)
    finally:
        workbook.close()


# ── Chunks ────────────────────────────────────────────────────────────────────

CHUNK_CHARS = 900
CHUNK_OVERLAP_LINES = 2


def chunk_text(text: str) -> list[str]:
    """Split text into ~CHUNK_CHARS pieces on line breaks, overlapping by a couple of lines."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    chunks, current, length = [], [], 0
    for line in lines:
        while len(line) > CHUNK_CHARS:  # one very long line (minified code, PDF paragraphs)
            if current:
                chunks.append("\n".join(current))
                current, length = [], 0
            chunks.append(line[:CHUNK_CHARS])
            line = line[CHUNK_CHARS:]
        if length + len(line) > CHUNK_CHARS and current:
            chunks.append("\n".join(current))
            current = current[-CHUNK_OVERLAP_LINES:]
            length = sum(len(l) for l in current)
        current.append(line)
        length += len(line)
    if current:
        chunks.append("\n".join(current))
    return chunks
