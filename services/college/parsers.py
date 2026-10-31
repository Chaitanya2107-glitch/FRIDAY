"""Bounded local timetable parsers; no paths, network, AI or persistence."""
import csv
from datetime import datetime, time
import io
from pathlib import PurePosixPath
import re
import zipfile

from services.college import models

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 300
EXTENSIONS = {'.png', '.jpg', '.jpeg', '.pdf', '.csv', '.docx', '.xlsx'}
DAYS = ('monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday')
HEADERS = {'day': {'day', 'weekday'}, 'start': {'start', 'starttime', 'from'},
           'end': {'end', 'endtime', 'to'}, 'subject': {'subject', 'course', 'class'}, 'room': {'room', 'venue'}}


def filename(name):
    if (not isinstance(name, str) or not name or len(name) > 160 or name in {'.', '..'}
            or re.search(r'[\\/:\x00-\x1f\x7f]', name) or name.rstrip(' .') != name
            or PurePosixPath(name).name != name):
        raise ValueError('Choose a file name without any path or control characters.')
    suffix = PurePosixPath(name).suffix.lower()
    if suffix not in EXTENSIONS:
        raise ValueError('Choose PNG, JPG, PDF, DOCX, XLSX or CSV.')
    return suffix


def _zip(data, extension):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 1500 or sum(i.file_size for i in members) > 40 * 1024 * 1024:
                raise ValueError('Office document expands beyond the safe import limit.')
            for item in members:
                parts = PurePosixPath(item.filename).parts
                if (item.flag_bits & 1 or '\\' in item.filename or item.filename.startswith('/')
                        or '..' in parts or ':' in item.filename or item.file_size > 20 * 1024 * 1024
                        or 'vbaproject' in item.filename.lower()):
                    raise ValueError('Unsupported or unsafe Office document.')
            names = {i.filename for i in members}
            required = 'word/document.xml' if extension == '.docx' else 'xl/workbook.xml'
            if not {'[Content_Types].xml', required} <= names:
                raise ValueError('The Office content does not match its extension.')
            for item in members:
                if item.filename.endswith(('.xml', '.rels')):
                    content = archive.read(item).upper()
                    if b'<!DOCTYPE' in content or b'<!ENTITY' in content:
                        raise ValueError('XML entities are not supported in uploads.')
    except zipfile.BadZipFile as error:
        raise ValueError('Invalid Office document.') from error


def validate(data, extension):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Choose a non-empty file no larger than 10 MB.')
    if extension in {'.docx', '.xlsx'}:
        if not data.startswith(b'PK\x03\x04'):
            raise ValueError('The file content does not match its extension.')
        _zip(data, extension)
    elif extension in {'.png', '.jpg', '.jpeg'}:
        from PIL import Image
        expected = b'\x89PNG\r\n\x1a\n' if extension == '.png' else b'\xff\xd8\xff'
        if not data.startswith(expected):
            raise ValueError('The image signature does not match its extension.')
        try:
            with Image.open(io.BytesIO(data)) as image:
                if image.width * image.height > 30_000_000:
                    raise ValueError('Image exceeds the 30 megapixel import limit.')
                if image.format != ('PNG' if extension == '.png' else 'JPEG'):
                    raise ValueError('Image format does not match its extension.')
                image.verify()
        except (OSError, Image.DecompressionBombError) as error:
            raise ValueError('Invalid or oversized image.') from error
    elif extension == '.pdf':
        from pypdf import PdfReader
        if not data.startswith(b'%PDF-'):
            raise ValueError('The PDF signature does not match its extension.')
        try:
            document = PdfReader(io.BytesIO(data))
            if document.is_encrypted or not document.pages or len(document.pages) > 30:
                raise ValueError('Use an unencrypted PDF with 1–30 pages (first three are imported).')
        except Exception as error:
            raise ValueError('Invalid, encrypted or oversized PDF.') from error
    elif extension == '.csv':
        text = data.decode('utf-8-sig')
        if any(ord(c) < 32 and c not in '\r\n\t' for c in text) or '\x7f' in text:
            raise ValueError('CSV must be UTF-8 text without binary/control bytes.')


def _cell(value):
    if isinstance(value, (datetime, time)):
        return value.strftime('%H:%M')
    return str(value).strip() if value is not None else ''


def _clock(value):
    raw = _cell(value)
    for fmt in ('%H:%M', '%H:%M:%S', '%I:%M %p', '%I %p', '%I:%M%p'):
        try:
            return datetime.strptime(raw.upper(), fmt).strftime('%H:%M')
        except ValueError:
            pass
    return raw


def _day(value):
    raw = _cell(value).lower()
    return next((d for d in DAYS if raw and d.startswith(raw[:3])), raw)


def from_table(table):
    bounded = []
    for index, row in enumerate(table):
        if index > MAX_ROWS or len(row) > 100:
            raise ValueError('The table exceeds 300 rows or 100 columns.')
        if any(_cell(c) for c in row):
            bounded.append(list(row))
    table = bounded
    if not table:
        return []
    if len(table) > MAX_ROWS + 1 or any(len(row) > 100 for row in table):
        raise ValueError('The table exceeds 300 rows or 100 columns.')
    header = [re.sub(r'[^a-z]', '', _cell(v).lower()) for v in table[0]]
    columns = {key: next((i for i, name in enumerate(header) if name in values), None) for key, values in HEADERS.items()}
    rows = []
    if columns['subject'] is not None:
        for raw in table[1:]:
            entry = {key: _cell(raw[i]) if i is not None and i < len(raw) else '' for key, i in columns.items()}
            entry.update(day=_day(entry['day']), start=_clock(entry['start']), end=_clock(entry['end']), room=entry['room'] or None)
            rows.append(entry)
    else:
        # Common day-column matrix: Time | Monday | Tuesday | ...
        days = {i: _day(name) for i, name in enumerate(header) if _day(name) in DAYS}
        if not days:
            return []
        for raw in table[1:]:
            times = re.split(r'\s*[-–—]\s*|\s+to\s+', _cell(raw[0]), maxsplit=1)
            for i, day in days.items():
                subject = _cell(raw[i]) if i < len(raw) else ''
                if subject and subject.lower() not in {'break', 'lunch', '-', 'free'}:
                    rows.append({'day': day, 'start': _clock(times[0]), 'end': _clock(times[1]) if len(times) == 2 else '',
                                 'subject': subject, 'room': None})
    if len(rows) > MAX_ROWS:
        raise ValueError('Import at most 300 class rows at a time.')
    return rows


def delimited(text):
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return from_table(csv.reader(io.StringIO(text), dialect))


def parse(data, extension):
    if extension == '.csv':
        rows = delimited(data.decode('utf-8-sig'))
    elif extension == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheet = workbook.active
            if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 100:
                raise ValueError('Use an active worksheet with at most 300 class rows and 100 columns.')
            rows = from_table(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    elif extension == '.docx':
        from docx import Document
        document = Document(io.BytesIO(data))
        rows = []
        for table in document.tables:
            rows.extend(from_table([[c.text for c in row.cells] for row in table.rows]))
        if not rows:
            rows = delimited('\n'.join(p.text for p in document.paragraphs))
    elif extension == '.pdf':
        from pypdf import PdfReader
        document = PdfReader(io.BytesIO(data))
        text = '\n'.join((page.extract_text() or '')[:150_000] for page in list(document.pages)[:3])
        rows = delimited(text)
        # An unfamiliar text layout remains local and editable. Only a PDF with
        # no extractable text may enter the explicitly consented vision path.
        return rows, not bool(text.strip())
    else:
        return [], True
    if not rows:
        raise ValueError('No class table found. Use Day, Start, End, Subject, Room column headings.')
    if len(rows) > MAX_ROWS:
        raise ValueError('Import at most 300 class rows at a time.')
    return rows, False


def warnings(rows, subjects):
    results = []
    for index, row in enumerate(rows):
        messages = []
        try:
            models.Slot.model_validate(row)
        except ValueError:
            messages.append('Missing or invalid day, subject or time')
        if isinstance(row.get('subject'), str) and row['subject'] and row['subject'].casefold() not in {s.casefold() for s in subjects}:
            messages.append('Unknown subject; check spelling or add subject details')
        text = lambda r, key: r.get(key, '') if isinstance(r.get(key, ''), str) else ''
        if any(text(row, 'day') == text(other, 'day') and text(row, 'start') < text(other, 'end')
               and text(other, 'start') < text(row, 'end') for other in rows[:index]):
            messages.append('Overlaps an earlier preview row')
        results.append(messages)
    return results


def images(data, extension):
    from PIL import Image
    from services.proactive import timetable
    result = []
    if extension != '.pdf':
        source = [Image.open(io.BytesIO(data)).convert('RGB')]
        document = None
    else:
        import pypdfium2
        document = pypdfium2.PdfDocument(data)
        source = []
    try:
        if document is not None:
            for i in range(min(len(document), timetable.MAX_PDF_PAGES)):
                page = document[i]
                try:
                    scale = min(2, timetable.MAX_IMAGE_SIDE / max(page.get_size()))
                    bitmap = page.render(scale=scale)
                    try:
                        source.append(bitmap.to_pil().convert('RGB'))
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        for image in source:
            image.thumbnail((timetable.MAX_IMAGE_SIDE, timetable.MAX_IMAGE_SIDE))
            output = io.BytesIO()
            image.save(output, format='JPEG', quality=90)
            result.append(output.getvalue())
            image.close()
        return result
    finally:
        if document is not None:
            document.close()
