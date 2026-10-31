"""Actual parser fixtures and the preview/confirm boundary; all paths are temporary."""
import asyncio
import copy
import io
import json
import zipfile

import pytest

from core.config import config
from services.college import models, parsers, schedules, store, uploads
from services.proactive import timetable

CSV = b'Day,Start,End,Subject,Room\nTuesday,09:00,10:00,DBMS,B204\n'
ROW = {'day': 'tuesday', 'start': '09:00', 'end': '10:00', 'subject': 'DBMS', 'room': 'B204'}


def fixture_bytes(extension):
    if extension == '.csv':
        return CSV
    output = io.BytesIO()
    if extension == '.xlsx':
        from openpyxl import Workbook
        document = Workbook()
        document.active.append(['Day', 'Start', 'End', 'Subject', 'Room'])
        document.active.append(['Tuesday', '09:00', '10:00', 'DBMS', 'B204'])
        document.save(output)
        document.close()
    elif extension == '.docx':
        from docx import Document
        document = Document()
        table = document.add_table(rows=2, cols=5)
        for cells, values in zip(table.rows, [['Day', 'Start', 'End', 'Subject', 'Room'], ['Tuesday', '09:00', '10:00', 'DBMS', 'B204']]):
            for cell, value in zip(cells.cells, values):
                cell.text = value
        document.save(output)
    else:
        from PIL import Image
        Image.new('RGB', (16, 16), '#152136').save(output, format='PNG')
    return output.getvalue()


def pdf_bytes(lines):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(('BT /F1 10 Tf 20 760 Td ' + ' '.join('(' + line + ') Tj 0 -14 Td' for line in lines) + ' ET').encode())
    page[NameObject('/Contents')] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    writer.close()
    return output.getvalue()


@pytest.mark.parametrize('lines,count,vision', [
    (['Day,Start,End,Subject,Room', 'Tuesday,09:00,10:00,DBMS,B204'], 1, False),
    (['A timetable layout with text but no labeled rows'], 0, False),
    ([], 0, True),
])
def test_real_pdf_text_stays_local_and_only_empty_text_needs_vision(lines, count, vision):
    staged = uploads.stage('synthetic.pdf', pdf_bytes(lines))
    assert len(staged['rows']) == count
    assert staged['requires_vision'] is vision
    assert not store.path('uploads').exists()
    if lines and not count:
        assert 'stays local' in staged['notice']


@pytest.mark.parametrize('extension', ['.csv', '.xlsx', '.docx'])
def test_real_fixtures_edit_preview_confirm_and_delete(extension):
    staged = uploads.stage('synthetic' + extension, fixture_bytes(extension))
    assert staged['rows'] == [ROW]
    assert timetable.load() == []
    assert not store.path('uploads').exists()  # not even raw bytes before Confirm
    rows = [{**staged['rows'][0], 'subject': 'OS'}]
    body = models.ImportEdit(rows=rows, revision=staged['revision'])
    checked = uploads.preview(staged['token'], body)
    assert checked['valid']
    assert checked['warnings'][0]  # unknown subject is a warning, not a silent mapping
    assert timetable.load() == []
    result = uploads.confirm(staged['token'], models.Confirm(**body.model_dump(), review=checked['review']))
    assert timetable.load()[0]['subject'] == 'OS'
    assert len(result['uploads']['items']) == 1
    record = result['uploads']['items'][0]
    file = store.path('uploads') / (record['id'] + extension)
    assert file.read_bytes() == fixture_bytes(extension) if extension == '.csv' else file.is_file()
    deleted = uploads.delete(record['id'], result['uploads']['revision'])
    assert deleted['items'] == []
    assert not file.exists()
    assert timetable.load()[0]['subject'] == 'OS'


def test_changed_rows_require_new_review_and_stale_timetable_refuses_confirm():
    staged = uploads.stage('test.csv', CSV)
    body = models.ImportEdit(rows=[ROW], revision=staged['revision'])
    checked = uploads.preview(staged['token'], body)
    altered = models.Confirm(rows=[{**ROW, 'subject': 'OS'}], revision=body.revision, review=checked['review'])
    with pytest.raises(store.Conflict, match='diff again'):
        uploads.confirm(staged['token'], altered)
    schedules.change(None, {**ROW, 'day': 'friday'}, schedules.snapshot()['revision'])
    with pytest.raises(store.Conflict, match='Changed elsewhere'):
        uploads.confirm(staged['token'], models.Confirm(**body.model_dump(), review=checked['review']))
    assert uploads.listing()['items'] == []


def test_review_refuses_overlaps_and_malformed_cells_without_crashing():
    staged = uploads.stage('test.csv', CSV)
    for bad in ({**ROW, 'subject': 12, 'start': {'bad': True}}, {**ROW, 'day': ''}):
        result = uploads.preview(staged['token'], models.ImportEdit(rows=[bad], revision=staged['revision']))
        assert not result['valid']
        assert result['warnings'][0]
    rows = [ROW, {**ROW, 'start': '09:30', 'end': '10:30'}]
    result = uploads.preview(staged['token'], models.ImportEdit(rows=rows, revision=staged['revision']))
    assert not result['valid'] and 'Overlaps' in ' '.join(result['warnings'][1])


@pytest.mark.parametrize('filename,data', [('fake.png', CSV), ('fake.jpg', b'\xff\xd8\xffgarbage'),
    ('fake.pdf', CSV), ('fake.xlsx', CSV), ('fake.docx', fixture_bytes('.xlsx')),
    ('binary.csv', b'Day,Subject\n\x00\x01'), ('script.exe', CSV), ('../escape.csv', CSV),
    ('C:\\temp\\table.csv', CSV), ('table.csv:stream', CSV), ('test.csv ', CSV)])
def test_type_spoofing_and_filename_paths_are_refused(filename, data):
    with pytest.raises(ValueError):
        uploads.stage(filename, data)
    assert timetable.load() == []


def test_oversize_and_zip_traversal_refused():
    with pytest.raises(ValueError, match='10 MB'):
        uploads.stage('large.csv', b'x' * (parsers.MAX_BYTES + 1))
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml', '<document/>')
        archive.writestr('../escape', 'bad')
    with pytest.raises(ValueError, match='unsafe'):
        uploads.stage('bad.docx', output.getvalue())


def test_csv_stops_consuming_rows_at_limit():
    consumed = []
    def rows():
        for i in range(10_000):
            consumed.append(i)
            yield ['Day', 'Start', 'End', 'Subject']
    with pytest.raises(ValueError, match='300 rows'):
        parsers.from_table(rows())
    assert len(consumed) < 304


def test_csv_day_column_matrix_is_local():
    rows, needs_vision = parsers.parse(b'Time,Monday,Tuesday\n09:00-10:00,DBMS,Lunch\n', '.csv')
    assert rows == [{'day': 'monday', 'start': '09:00', 'end': '10:00', 'subject': 'DBMS', 'room': None}]
    assert needs_vision is False


def test_stage_is_bounded_expiring_and_discardable(monkeypatch):
    items = [uploads.stage(f'test{i}.csv', CSV) for i in range(4)]
    with pytest.raises(ValueError, match='Four previews'):
        uploads.stage('fifth.csv', CSV)
    uploads.discard(items[0]['token'])
    assert not store.path('uploads').exists()
    with pytest.raises(store.Conflict, match='expired'):
        uploads.preview(items[0]['token'], models.ImportEdit(rows=[ROW], revision=items[0]['revision']))
    monkeypatch.setattr(uploads.time, 'monotonic', lambda: 10 ** 20)
    with pytest.raises(store.Conflict, match='expired'):
        uploads.preview(items[1]['token'], models.ImportEdit(rows=[ROW], revision=items[1]['revision']))


def test_image_requires_fresh_consent_and_policy_is_preserved(monkeypatch):
    staged = uploads.stage('test.png', fixture_bytes('.png'))
    assert staged['requires_vision'] and not staged['rows']
    monkeypatch.setattr(config, 'AI_ASSIST_ONLY', True)
    monkeypatch.setattr('providers.failover.provider_chain', lambda primary: ['ollama'])
    monkeypatch.setattr(config, 'VISION_PROVIDER', 'ollama')
    monkeypatch.setattr(config, 'VISION_MODEL', 'gemma3:4b')
    state = uploads.vision()
    assert not state['available']
    with pytest.raises(ValueError):
        asyncio.run(uploads.extract(staged['token'], state['consent']))
    monkeypatch.setattr(config, 'AI_ASSIST_ONLY', False)
    current = uploads.vision()
    called = []
    async def extract(images, provider_name, section='', approved_model=None):
        called.append((provider_name, approved_model))
        return [ROW]
    monkeypatch.setattr(timetable, 'extract_classes', extract)
    with pytest.raises(store.Conflict, match='consent again'):
        asyncio.run(uploads.extract(staged['token'], state['consent']))
    assert called == []
    result = asyncio.run(uploads.extract(staged['token'], current['consent']))
    assert result['rows'] == [ROW] and called == [('ollama', 'gemma3:4b')]
    assert not store.path('uploads').exists()


def test_openrouter_upload_uses_only_the_consented_model(monkeypatch):
    from types import SimpleNamespace
    from providers.base_provider import ProviderResponse
    calls = []
    async def fixed(model, messages, max_tokens, temperature, tools):
        calls.append(model)
        return ProviderResponse(content=json.dumps([ROW]), provider='openrouter', model=model)
    async def forbidden(**kwargs):
        pytest.fail('The provider-wide fallback path must not be used by consented uploads')
    monkeypatch.setattr('providers.registry.get_provider', lambda *a, **kw: SimpleNamespace(_complete_with=fixed, complete=forbidden))
    result = asyncio.run(timetable.extract_classes([b'synthetic jpeg'], 'openrouter', approved_model='approved-model'))
    assert calls == ['approved-model']
    assert result[0]['subject'] == 'DBMS'


def test_failure_during_confirm_rolls_back_schedule_manifest_and_original(monkeypatch):
    from services import settings_store
    staged = uploads.stage('test.csv', CSV)
    body = models.ImportEdit(rows=[ROW], revision=staged['revision'])
    checked = uploads.preview(staged['token'], body)
    publish = settings_store.publish
    def fail_schedule(path, data):
        if path == timetable.path:
            raise OSError('Synthetic schedule publication failure')
        return publish(path, data)
    monkeypatch.setattr(settings_store, 'publish', fail_schedule)
    with pytest.raises(OSError):
        uploads.confirm(staged['token'], models.Confirm(**body.model_dump(), review=checked['review']))
    assert timetable.load() == []
    assert uploads.listing()['items'] == []
    assert not list(store.path('uploads').glob('*.csv'))
